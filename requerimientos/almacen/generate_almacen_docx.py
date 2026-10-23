import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_almacen_docx():
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

    # Colors
    NAVY = RGBColor(0x0f, 0x17, 0x2a) # Slate 900
    INDIGO = RGBColor(0x37, 0x30, 0xa3) # Indigo 800
    GRAY_TEXT = RGBColor(0x64, 0x74, 0x8b) # Slate 500
    TEAL = RGBColor(0x0f, 0x76, 0x6e) # Teal 700

    # Document Header / Pre-title
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_after = Pt(2)
    run_pre = p_pre.add_run("FORMATEX WMS 360° · ESPECIFICACIÓN DE REGLAS DE NEGOCIO Y OPERACIÓN")
    run_pre.font.size = Pt(9)
    run_pre.font.bold = True
    run_pre.font.color.rgb = INDIGO

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Módulo de Almacén, Surtido, Corte, Empaque, Embarques y Telares")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    # Subtitle / Version Badge
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Versión 1.0 Oficial Integral · Levantamiento Operativo en Bodega (23 de Octubre de 2026)")
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = GRAY_TEXT

    # Metadata Table Box
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    data_meta = [
        ("Organización Solicitante:", "Forma Textil S. de R.L. de C.V. (Formatex)"),
        ("Socio Tecnológico:", "Movida TCI LLC"),
        ("Fecha de Levantamiento:", "Viernes, 23 de Octubre de 2026"),
        ("Participantes de Formatex:", "Israel (Jefe de Almacén), Andrea Ramírez (ATC/Operaciones), Santiago (Comercial)"),
        ("Participantes de Movida TCI:", "Ricardo Zerpa (Arquitectura de Procesos), Jonathan Palacios (Dirección Técnica)")
    ]

    for idx, (lbl, val) in enumerate(data_meta):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.5)
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, 70, 70, 100, 100)
        set_cell_margins(cell_val, 70, 70, 100, 100)

        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = NAVY

        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helpers
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = INDIGO
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.bold = True
        r_pre.font.color.rgb = NAVY
        p.add_run(text)
        return p

    def add_callout(text, title="NOTA OPERATIVA VINCULANTE:"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.7)
        set_cell_background(cell, "EFF6FF")
        set_cell_margins(cell, 100, 100, 140, 140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"{title} ")
        r_title.font.bold = True
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = INDIGO
        r_text = p.add_run(text)
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Propósito
    add_h1("1. Propósito y Marco de Aceptación Vinculante")
    p = doc.add_paragraph(
        "El presente documento recopila, formaliza y desglosa minuciosamente cada una de las reglas operativas, "
        "fórmulas matemáticas, tiempos de promesa (SLA), capacidades de producción y flujos interdepartamentales "
        "consensuados con la jefatura de Almacén de Formatex (Israel) y el equipo operativo."
    )
    add_callout(
        "Este documento actúa como anexo técnico vinculante del sistema WMS 360°. Establece con precisión matemática y quirúrgica "
        "la interacción entre Atención al Cliente (ATC), Almacén, Mesa de Corte, Empaque, Embarques, Reparto Propio, "
        "Fleteras Externas y Reabastecimiento a Telares, eliminando ambigüedades operativas.",
        "FINALIDAD OPERATIVA Y CONTRACTUAL:"
    )

    # 2. Regla 1: Protocolo <100m
    add_h1("2. Regla 1: Protocolo de Cotización y Validación de Existencia (< 100 Metros)")
    add_h2("2.1. El Umbral Crítico de los 100 Metros")
    doc.add_paragraph(
        "En el inventario textil de Formatex, cuando un SKU cuenta con menos de 100 metros en existencia física disponible "
        "o cuando el pedido involucra rollos abiertos (remanentes), el sistema WMS bloquea la venta automática o apartado directo "
        "y dispara un protocolo de validación física obligatoria en piso de bodega."
    )
    add_bullet("Consulta de Existencia Rápida: ", "El almacenista acude al rack, revisa la etiqueta física y notifica a ATC en un SLA de 15 a 30 minutos (máximo 1 hora).")
    add_bullet("Cotización en Firme: ", "Si el cliente está en proceso de pago/cierre, el rollo se baja del rack y se traslada a la mesa de medición para verificar el metraje metro a metro, descartando variaciones por elasticidad o mermas.")

    # 3. Regla 2: Apartados
    add_h1("3. Regla 2: Políticas de Apartado y Retención de Inventario (2 Días Cliente vs. 3 Días Almacén)")
    doc.add_paragraph("Para sincronizar la presión comercial y la realidad administrativa de pagos:")
    add_bullet("Ventana Comercial Externa (Cliente): ", "2 días hábiles (48 horas) comunicados al cliente en cotización para acelerar el depósito bancario.")
    add_bullet("Ventana Operativa Interna (Almacén): ", "3 días hábiles (72 horas) de resguardo físico en bodega, protegiendo al cliente ante ciclos de pago de fin de semana o viernes.")
    add_bullet("Marbete Físico en Bodega: ", "Almacén imprime y coloca el marbete 'APARTADO - PROHIBIDA SU VENTA' con pedido, cliente y fecha límite.")
    add_bullet("Desapartado Automático: ", "Al concluir el 3er día hábil sin comprobante de pago, el WMS libera automáticamente el inventario a DISPONIBLE y emite alerta en la Zebra para retirar el marbete físico.")

    # 4. Regla 3: Capacidad y Productividad
    add_h1("4. Regla 3: Capacidad Operativa y Asignación de Surtido por Altura y Especialidad")
    doc.add_paragraph(
        "Se homologa la métrica de productividad oficial: cada surtidor procesa entre 55 y 60 unidades (cortes o rollos cerrados) por jornada laboral completa (08:00 a 17:30 / 18:00 hrs)."
    )
    add_bullet("Rollos en Racks Altos: ", "Niveles superiores. Asignación exclusiva a personal con licencia y certificación de Montacargas.")
    add_bullet("Cortes a Nivel de Piso: ", "Rollos abiertos en niveles bajos. Asignables a personal de nuevo ingreso o general.")
    add_bullet("Pedidos Mixtos (Cortes en Piso + Rollos en Altura): ", "Regla de oro: Asignados íntegramente a un surtidor experto con montacargas. Queda estrictamente prohibido dividir el pedido para evitar que los cortes queden varados esperando los rollos de altura.")

    # 5. Regla 4: Carritos A1, A2, A3
    add_h1("5. Regla 4: Flujo de Surtido, Carritos de Corte de 3 Niveles (A1, A2, A3) y Control de Calidad")
    doc.add_paragraph(
        "El surtidor traslada los rollos al área de corte en carritos metálicos organizados en 3 niveles:"
    )
    add_bullet("Estructura de Carritos: ", "Bandeja Nivel 1 (A1), Bandeja Nivel 2 (A2) y Bandeja Nivel 3 (A3).")
    add_bullet("Registro en WMS: ", "El surtidor registra en la Zebra el carrito y nivel exacto donde coloca el rollo (ej. Carrito 02 - Nivel A1). El cortador visualiza esta ubicación al instante.")
    add_bullet("Control de Calidad en Mesa de Corte: ", "Al extender el rollo, si se detecta defecto de fábrica, barra o mancha, se suspende el corte, se registra la incidencia en el WMS y se solicita cambio de rollo del mismo lote al surtidor. Si no hay más rollos del mismo lote, se alerta de inmediato a ATC.")

    # 6. Regla 5: SLA y Regla Anti-Trampa
    add_h1("6. Regla 5: Matriz Oficial de Tiempos de Promesa (SLA) y Regla Anti-Trampa")
    
    # SLA Table
    sla_table = doc.add_table(rows=5, cols=3)
    sla_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sla_table.autofit = False

    sla_data = [
        ("Tipo de Pedido", "Composición / Cantidad", "Tiempo Promesa Oficial (SLA)"),
        ("Rollo Cerrado Estándar", "Menos de 10 rollos cerrados", "24 a 48 Horas Hábiles"),
        ("Cortes Menores", "1 a 4 cortes", "72 Horas Hábiles (3 días)"),
        ("Cortes Medios", "5 a 9 cortes", "96 Horas Hábiles (4 días)"),
        ("Cortes Mayores / Masivos", "10 cortes o más", "96+ Horas Hábiles (Con Almacén)")
    ]

    for idx, row in enumerate(sla_table.rows):
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(2.2)
        c1.width = Inches(2.3)
        c2.width = Inches(2.2)
        bg = "0F172A" if idx == 0 else ("F8FAFC" if idx % 2 == 1 else "FFFFFF")
        fg = RGBColor(0xFF, 0xFF, 0xFF) if idx == 0 else RGBColor(0x1E, 0x29, 0x3B)
        for i, cell in enumerate([c0, c1, c2]):
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(sla_data[idx][i])
            r.font.bold = (idx == 0 or i == 2)
            r.font.size = Pt(9.5)
            r.font.color.rgb = fg

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        "Regla Anti-Trampa: Si un mismo cliente fiscal divide un pedido de cortes en varias órdenes el mismo día "
        "(ej. dos pedidos de 4 cortes en lugar de uno de 8 para forzar un SLA de 72h), el WMS consolida automáticamente "
        "el número total de cortes y asigna el SLA global correspondiente (8 cortes = 96 horas hábiles).",
        "PROTECCIÓN ANTI-TRAMPA DE SLA:"
    )

    # 7. Regla 6: Lotes y Variación de Tono
    add_h1("7. Regla 6: Algoritmo de Sugerencia por Lote y Alerta Preventiva de Tono")
    doc.add_paragraph(
        "El motor de asignación prioriza rigurosamente abastecer pedidos utilizando rollos del mismo lote de tintada. "
        "Cuando el stock exige combinar lotes diferentes, se activa un bloqueo preventivo:"
    )
    add_bullet("Alerta Bloqueante en ATC: ", "El sistema despliega una advertencia obligatoria de 'Posible Variación de Tono por Mezcla de Lotes'.")
    add_bullet("Validación Comercial: ", "ATC debe contactar al cliente para validar el uso: si es para una misma pieza/sala modular, la mezcla queda prohibida; si es para piezas separadas (cojines/sillas), el cliente puede autorizar la entrega.")

    # 8. Regla 7: Las 4 Etiquetas de Embarque
    add_h1("8. Regla 7: Las 4 Etiquetas Oficiales de Embarque y Estiba por Niveles")
    
    lbl_table = doc.add_table(rows=5, cols=4)
    lbl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lbl_table.autofit = False

    lbl_data = [
        ("Tipo de Etiqueta", "Destino / Canal", "Ubicación en Embarque", "Particularidad Operativa"),
        ("1. Local", "Reparto ZMG en camiones propios Formatex", "Nivel Inferior (Piso)", "Salida matutina en rutas ZMG."),
        ("2. Foránea", "Fleteras foráneas y fletes pesados", "Nivel Superior (Racks)", "Esperan recolección de transportistas."),
        ("3. Ocotlán", "Ruta fija Ocotlán / Poncitlán", "Área de Ruta Ocotlán", "Salida fija Martes y Jueves mañana."),
        ("4. Paqueterías", "Paquetexpress, Travisa, Tres Guerras", "Andén de Paquetería", "Corte diario a las 14:00 hrs.")
    ]

    for idx, row in enumerate(lbl_table.rows):
        c0, c1, c2, c3 = row.cells[0], row.cells[1], row.cells[2], row.cells[3]
        c0.width = Inches(1.4)
        c1.width = Inches(1.8)
        c2.width = Inches(1.7)
        c3.width = Inches(1.8)
        bg = "0F172A" if idx == 0 else ("F8FAFC" if idx % 2 == 1 else "FFFFFF")
        fg = RGBColor(0xFF, 0xFF, 0xFF) if idx == 0 else RGBColor(0x1E, 0x29, 0x3B)
        for i, cell in enumerate([c0, c1, c2, c3]):
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(lbl_data[idx][i])
            r.font.bold = (idx == 0 or i == 0)
            r.font.size = Pt(9.5)
            r.font.color.rgb = fg

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_bullet("Requerimiento Paquetexpress: ", "Empaque captura peso real en kg y medidas volumétricas (L x A x A cm) para emisión de pre-guía.")

    # 9. Regla 8: Calendario de Rutas
    add_h1("9. Regla 8: Calendario y Horarios de Salida de Rutas y Paqueterías")
    add_bullet("Ruta Ocotlán: ", "Salida Martes y Jueves a primera hora (surtido y empaque Lunes y Miércoles previos antes de 18:00 hrs).")
    add_bullet("Ruta Arandas / Altos: ", "Salida Miércoles a primera hora (surtido Martes).")
    add_bullet("Paqueterías Regulares: ", "Corte operativo diario a las 14:00 hrs. Pedidos ingresados después de las 14:00 hrs se procesan para recolección del día hábil siguiente.")

    # 10. Regla 9: Evidencia Digital de Guías
    add_h1("10. Regla 9: Evidencia Digital de Guías de Fletera y Chofer (Tracking)")
    doc.add_paragraph(
        "Se digitaliza la captura de acuses de recibo y guías para eliminar fotos sueltas en grupos de WhatsApp:"
    )
    add_bullet("Captura de Guía y Fletera: ", "El chofer o almacenista ingresa en el WMS el folio de tracking oficial y la fletera.")
    add_bullet("Fotografía Obligatoria del Acuse: ", "Se captura con la cámara del dispositivo móvil la guía sellada/firmada.")
    add_bullet("Actualización Automática: ", "El pedido pasa a EN_TRANSITO, adjuntando la fotografía al expediente del pedido y notificando a ATC y al cliente con su enlace de rastreo.")

    # 11. Regla 10: Reabastecimiento a Telares
    add_h1("11. Regla 10: Reabastecimiento Formatex ➔ Telares (Suministro Máster de 8 Meses)")
    add_bullet("Meta de Stock Formatex: ", "Israel establece una meta de 8 meses de inventario (el doble de los 4 meses calculados en el máster de compras comercial) para absorber demoras de contenedores marítimos.")
    add_bullet("Requisición Ordinaria: ", "Corridas periódicas los Lunes y Viernes basadas en la fórmula de 8 meses vs stock, pedidos y tránsito.")
    add_bullet("Corte de Pedidos Especiales ATC: ", "Corte diario a las 13:00 hrs; envío consolidado a Telares a las 14:00 hrs; surtido al día siguiente.")
    add_bullet("Doble Etiquetado: ", "Telares etiqueta el contenedor de China; Formatex coloca su propia etiqueta con SKU, lote y ubicación asignada al recibir en nave.")

    # Signatures Table
    add_h1("12. Firmas de Conformidad y Aceptación Operativa")
    sig_table = doc.add_table(rows=3, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    sig_headers = [
        ("Por Forma Textil S. de R.L. de C.V.", "Por Movida TCI LLC"),
        ("Israel\nJefe de Almacén & Logística\n\nAndrea Ramírez\nAtención al Cliente & Operaciones\n\nSantiago\nOperaciones & Comercial", 
         "Jonathan Palacios\nDirector de Tecnología\n\nRicardo Zerpa\nLíder de Producto & Arquitectura\n\nFecha: 23 de Octubre de 2026"),
        ("Firma de Conformidad: ______________________", "Firma de Conformidad: ______________________")
    ]

    for idx, row in enumerate(sig_table.rows):
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(3.3)
        c1.width = Inches(3.4)
        bg = "0F172A" if idx == 0 else ("F8FAFC" if idx == 1 else "FFFFFF")
        fg = RGBColor(0xFF, 0xFF, 0xFF) if idx == 0 else RGBColor(0x1E, 0x29, 0x3B)
        for i, cell in enumerate([c0, c1]):
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(sig_headers[idx][i])
            r.font.bold = (idx == 0 or idx == 2)
            r.font.size = Pt(9.5)
            r.font.color.rgb = fg

    # Output path
    out_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(out_dir, "REGLAS_DE_NEGOCIO_ALMACEN_FORMATEX.docx")
    doc.save(out_path)
    print(f"Documento DOCX generado exitosamente en: {out_path}")

if __name__ == "__main__":
    generate_almacen_docx()
