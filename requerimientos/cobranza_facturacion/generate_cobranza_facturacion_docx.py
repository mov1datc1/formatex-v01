import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="F8FAFC", border_hex="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.7)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.font.bold = True
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_cobranza_docx():
    doc = docx.Document()

    # Set page margins (0.8 in top/bottom, 0.85 in left/right)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    NAVY = RGBColor(0x0f, 0x17, 0x2a)      # Slate 900
    GOLD = RGBColor(0xd9, 0x77, 0x06)      # Amber 600
    INDIGO = RGBColor(0x37, 0x30, 0xa3)    # Indigo 800
    EMERALD = RGBColor(0x05, 0x96, 0x69)   # Emerald 600
    GRAY_TEXT = RGBColor(0x64, 0x74, 0x8b) # Slate 500

    # Document Header / Pre-title
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_after = Pt(2)
    run_pre = p_pre.add_run("FORMATEX WMS 360° · ESPECIFICACIÓN DE REGLAS DE NEGOCIO Y OPERACIÓN")
    run_pre.font.size = Pt(9)
    run_pre.font.bold = True
    run_pre.font.color.rgb = INDIGO

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Módulo de Finanzas, Crédito, Cobranza, Facturación & Comisiones")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Versión 1.0 Oficial Consolidada · Sesión de Trabajo con Dirección de Finanzas (26 de Octubre de 2026)")
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = GRAY_TEXT

    # Metadata Table Box
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    data_meta = [
        ("Organización Solicitante:", "Forma Textil S. de R.L. de C.V. (Formatex - Guadalajara, Jalisco)"),
        ("Socio Tecnológico:", "Movida TCI LLC"),
        ("Fecha de Levantamiento:", "Lunes, 26 de Octubre de 2026"),
        ("Participantes Formatex:", "César (Director de Finanzas, Crédito y Cobranza), Andrea Ramírez (ATC / Operaciones), Santiago (Comercial), Mayra García (Administración)"),
        ("Participantes Movida TCI:", "Jonathan Palacios (Dirección Técnica / Ingeniería WMS 360+)")
    ]

    for idx, (lbl, val) in enumerate(data_meta):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.5)
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)

        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(lbl)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = NAVY

        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper function for headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = INDIGO
        return p

    # SECTION 1: RESUMEN EJECUTIVO
    add_h1("1. Resumen Ejecutivo y Propósito")
    p = doc.add_paragraph(
        "El presente documento formaliza las reglas de negocio, directrices operativas y tratamientos fiscales "
        "establecidos en la sesión de levantamiento con César (Director de Finanzas, Crédito y Cobranza) "
        "para el sistema Formatex WMS 360°. Este módulo representa el engranaje crítico entre la captura comercial "
        "(ATC/Ventas), la operación física del almacén (Surtido/Corte/Empaque) y la certeza patrimonial ante "
        "el Servicio de Administración Tributaria (SAT)."
    )
    p.paragraph_format.space_after = Pt(6)

    add_callout(
        doc,
        "Principio Rector de la Dirección Financiera",
        "Ninguna mercancía sale de rampa si no cuenta con respaldo bancario confirmado (en contado) o con autorización "
        "explícita y negociada de Cobranza (en crédito). Las comisiones de ventas operan bajo la estricta máxima: "
        "'Mes trabajado = Mes cobrado'. No se comisiona ni un centavo sobre facturas con saldos insolutos pendientes.",
        bg_hex="EFF6FF",
        border_hex="2563EB"
    )

    # SECTION 2: POLÍTICA DE COBRANZA Y LIBERACIÓN DE PEDIDOS
    add_h1("2. Política de Cobranza y Liberación de Pedidos")
    
    add_h2("2.1 Clientes de Contado (Verificación Bancaria Obligatoria)")
    p = doc.add_paragraph(
        "Para los pedidos de clientes bajo condición de Contado, el pedido nace formalmente en almacén hasta que el área "
        "de Cobranza verifica fehacientemente que el depósito o transferencia interbancaria se encuentra reflejado en las "
        "cuentas fiscales de Formatex (BBVA Bancomer / Santander)."
    )
    doc.add_paragraph(
        "1. Atención a Clientes (ATC) genera la cotización y recibe el comprobante por WhatsApp/Email.\n"
        "2. El pedido entra a la cola de Cobranza en estado PAGO_RECIBIDO (no visible para picker ni almacén).\n"
        "3. El analista de cobranza valida el estado de cuenta bancario, corrobora la referencia y da clic en 'Validar Pago'.\n"
        "4. El sistema registra cobradorNombre, fechaAprobCobranza y cambia el estado a POR_SURTIR, liberando el pedido a la Torre de Despacho de Almacén.",
        style='List Bullet'
    )

    add_h2("2.2 Clientes de Crédito y Manejo de Pedidos Retenidos (RETENIDO_CREDITO)")
    p = doc.add_paragraph(
        "Para clientes con línea de crédito autorizada, los plazos comerciales oficiales son de: 7, 15, 30, 45 y 60 días naturales. "
        "En la práctica comercial, los clientes tienden a exceder los plazos pactados (el de 30 días paga a 45 días, el de 60 paga a 90 días). "
        "Para mitigar este riesgo sin frenar la labor de venta, César definió una regla de negocio fundamental:"
    )
    
    add_callout(
        doc,
        "Regla de Captura Permitida pero Retención Automática",
        "El sistema NUNCA debe bloquear la captura del pedido en ATC. El cliente o agente puede registrar el pedido para que 'se ilusionen' "
        "y se mantenga el registro de su necesidad. Sin embargo, si el cliente presenta facturas vencidas o sobregiro de línea, "
        "el pedido pasa automáticamente al estado RETENIDO_CREDITO.\n\n"
        "El pedido retenido NO pasa a surtido ni aparta tela indefinidamente. Se mantiene en espera de que Cobranza negocie la recuperación de cartera.",
        bg_hex="FEF3C7",
        border_hex="D97706"
    )

    add_h2("2.3 Criterios de Negociación Discrecional de Cobranza")
    p = doc.add_paragraph(
        "La liberación de un pedido retenido es una facultad estrictamente discrecional de César y su equipo. No se somete a "
        "una fórmula matemática rígida, ya que depende del valor estratégico del cliente, la antigüedad de su saldo y el tipo de tela. "
        "Ejemplos de acuerdos típicos en la operación:\n"
        "• 'El cliente debe 4 facturas por $60,000 MXN. Mete pedido urgente de $15,000 MXN. Se negocia que deposite $25,000 MXN a facturas viejas para liberarle el pedido nuevo.'\n"
        "• 'Cliente con factura vencida de hace 5 días por volumen habitual: Cobranza autoriza la liberación tras promesa de pago por escrito.'\n\n"
        "En el WMS 360°, el botón 'Liberar Pedido por Negociación' solicita obligatoriamente el ingreso de notas de negociación, "
        "registra el usuario responsable y transfiere el pedido a POR_SURTIR."
    )

    add_h2("2.4 Tolerancia Financiera de Redondeo (±$10.00 MXN)")
    p = doc.add_paragraph(
        "Por diferencias de comisiones bancarias, redondeos en cotizaciones de metros fraccionarios o ajustes en ventanilla, "
        "César cuenta con una tolerancia de ±$10.00 MXN. Si el saldo restante de una factura tras aplicar un abono es menor o igual a "
        "$10.00 MXN, el sistema debe permitir marcar la factura como PAGADA al 100%, absorbiendo la discrepancia sin alterar la póliza contable."
    )

    # SECTION 3: COBRANZA PREVENTIVA Y PLANTILLA DE CORREO (5 DÍAS ANTES)
    add_h1("3. Cobranza Preventiva Automatizada (Alerta a 5 Días)")
    p = doc.add_paragraph(
        "A propuesta consensuada entre Andrea Ramírez, Santiago y César, el sistema implementará un motor de cobranza preventiva. "
        "Exactamente cinco (5) días naturales antes de que expire la fecha límite de pago de una factura de crédito (PPD), "
        "el WMS enviará un correo automático personalizado al cliente con copia al agente de ventas asignado."
    )

    add_callout(
        doc,
        "Plantilla Oficial de Correo de Cobranza Preventiva Formatex",
        "ASUNTO: Recordatorio Preventivo: Tu Factura {{folioFactura}} está próxima a vencer - Formatex\n\n"
        "Estimado(a) {{nombreCliente}},\n\n"
        "Esperamos que te encuentres excelente. En Formatex valoramos enormemente tu preferencia y queremos apoyarte para mantener "
        "tu cuenta al corriente y asegurar que tus pedidos futuros y disponibilidad de telas se procesen con la máxima agilidad.\n\n"
        "Te recordamos que tu factura {{folioFactura}}, con un saldo pendiente de {{montoInsoluto}} MXN, tiene fecha límite de vencimiento el próximo {{fechaVencimiento}} (quedan {{diasRestantes}} días naturales).\n\n"
        "DATOS DE PAGO:\n"
        "• Beneficiario: FORMA TEXTIL S. DE R.L. DE C.V.\n"
        "• Banco: BBVA Bancomer | Cuenta: 0112345678 | CLABE: 012320001123456789\n"
        "• Referencia: {{codigoPedido}} / {{folioFactura}}\n\n"
        "Una vez realizado tu depósito o transferencia, te agradeceremos compartir el comprobante al correo cobranza@formatex.mx o a través de tu asesor comercial {{agenteVentas}} para aplicar tu pago de inmediato y emitir tu Complemento de Pago (REP).\n\n"
        "Agradecemos de antemano tu puntualidad y quedamos a tus órdenes.\n\n"
        "Atentamente,\n"
        "Departamento de Crédito y Cobranza\n"
        "Forma Textil S. de R.L. de C.V. | Guadalajara, Jalisco",
        bg_hex="F0FDF4",
        border_hex="16A34A"
    )

    # SECTION 4: FACTURACIÓN CFDI 4.0 Y TRATAMIENTO FISCAL
    add_h1("4. Facturación CFDI 4.0 y Tratamiento Fiscal (SAT / Facturapi)")
    
    add_h2("4.1 Momento Preciso de Emisión")
    p = doc.add_paragraph(
        "La factura fiscal CFDI 4.0 se emite exclusivamente cuando el pedido ha sido completamente surtido y empacado (estado EMPACADO). "
        "En este punto se conoce el metraje exacto empacado en rollos y cortes, eliminando al 100% las diferencias de centavos o metros "
        "entre cotización y factura."
    )

    add_h2("4.2 Facturas PUE (Contado) - Prohibición Absoluta de Complemento REP")
    p = doc.add_paragraph(
        "Conforme a las reglas del SAT (Anexo 20), las facturas emitidas bajo método PUE (Pago en una sola exhibición) amparan operaciones "
        "cuyo cobro se consumó antes o al momento de la emisión. César enfatizó que en CONTPAQi y en el WMS:\n"
        "• La aplicación del cobro de facturas PUE es puramente interna ('No Timbrada').\n"
        "• Queda ESTRICTAMENTE PROHIBIDO emitir un Recibo Electrónico de Pago (REP / CFDI tipo P) para facturas PUE, ya que constituiría una infracción fiscal."
    )

    add_h2("4.3 Facturas PPD (Crédito) y Complemento de Pago por Cada Abono")
    p = doc.add_paragraph(
        "Las ventas a crédito se facturan bajo método PPD (Pago en Parcialidades o Diferido) con forma de pago '99 Por Definir'. "
        "Para estos documentos rige la siguiente obligación fiscal inquebrantable:"
    )
    add_callout(
        doc,
        "Obligatoriedad de Complemento de Pago (REP) por Cada Abono Individual",
        "Cada transferencia o depósito bancario que un cliente a crédito realiza —incluso si se trata de un abono parcial de $5,000 MXN "
        "aplicado a una factura de $25,000 MXN— REQUIERE OBLIGATORIAMENTE la emisión de su propio Complemento de Pago (REP) timbrado "
        "ante el SAT a través de Facturapi, relacionando el UUID de la factura origen y especificando el saldo anterior, el importe pagado "
        "y el saldo insoluto restante.",
        bg_hex="FAF5FF",
        border_hex="9333EA"
    )

    add_h2("4.4 Notas de Crédito y Devoluciones con Relación CFDI")
    p = doc.add_paragraph(
        "Cualquier devolución de mercancía o ajuste por flete/garantía de calidad genera una Nota de Crédito (CFDI tipo Egreso) "
        "timbrada ante el SAT, la cual debe incluir el nodo de CFDIs Relacionados:\n"
        "• Tipo de Relación '01' (Nota de crédito de los documentos relacionados) o '03' (Devolución de mercancía sobre facturas o traslados previos).\n"
        "• UUID fiscal de la factura de origen vinculada."
    )

    # SECTION 5: REGLAMENTO OFICIAL DE COMISIONES DE VENTAS
    add_h1("5. Reglamento Oficial de Comisiones de Ventas")
    
    add_h2("5.1 Principio Rector: 'Mes Trabajado = Mes Cobrado'")
    p = doc.add_paragraph(
        "Las comisiones de ventas de Formatex no se generan por levantar cotizaciones ni por emitir facturas. "
        "El derecho a comisionar nace estrictamente cuando la factura asociada al pedido ha sido pagada al 100%. "
        "Si una factura tiene un saldo insoluto pendiente al cierre de mes, no genera comisión en ese periodo."
    )

    add_h2("5.2 Calendario y Canal de Dispersión: Día 10 de Cada Mes")
    p = doc.add_paragraph(
        "Las comisiones se liquidan de manera mensual los días 10 de cada mes, cubriendo las facturas cobradas en su totalidad "
        "durante el mes calendario inmediato anterior (ej. el 10 de octubre se liquidaron las facturas cobradas en septiembre). "
        "César remite a cada asesor su estado de cuenta individualizado por correo electrónico confidencial. Ningún vendedor puede ver "
        "los importes de sus colegas."
    )

    add_h2("5.3 Base de Cálculo Gravable e Impacto de Notas de Crédito")
    p = doc.add_paragraph(
        "• La comisión se calcula sobre el Monto Bruto (subtotal de tela antes de IVA y retenciones).\n"
        "• Cualquier devolución física de mercancía o nota de crédito por garantía/flete RESTA directamente la base comisionable.\n"
        "Ejemplo: Factura de $1,000 MXN. Devolución de $200 MXN + Nota de crédito por flete de $300 MXN = Base ajustada de $500 MXN. "
        "La comisión del agente se calcula exclusivamente sobre los $500 MXN netos cobrados."
    )

    add_h2("5.4 Cadena Jerárquica de Comisiones")
    
    table_com = doc.add_table(rows=4, cols=3)
    table_com.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_com.autofit = False

    headers_com = ["Rol Jerárquico", "Base de Cálculo", "Porcentaje / Regla"]
    for i, h in enumerate(headers_com):
        c = table_com.rows[0].cells[i]
        set_cell_background(c, "1E293B")
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)

    rows_com = [
        ("1. Agente / Vendedor de Campo", "Subtotal neto cobrado de clientes asignados a su cartera", "Variable según lista de precios (F1 a F10) asignada al pedido"),
        ("2. Gerente de Ventas", "Subtotal neto cobrado de todos los agentes que integran su equipo comercial", "Porcentaje menor fijo sobre la venta consolidada de su grupo"),
        ("3. Dirección Comercial", "Subtotal neto cobrado de todos los gerentes de ventas de Formatex", "Mismo porcentaje del gerente, pero multiplicado por los 3-4 gerentes a su cargo")
    ]

    for idx, (rol, base, pct) in enumerate(rows_com):
        row = table_com.rows[idx + 1]
        for c_idx, val in enumerate([rol, base, pct]):
            c = row.cells[c_idx]
            set_cell_background(c, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("5.5 Requisito Fiscal Indispensable del Vendedor")
    p = doc.add_paragraph(
        "Para que la empresa proceda a la dispersión bancaria de la comisión, el agente de ventas o gerente debe emitir y enviar "
        "previamente una factura fiscal (CFDI de ingresos bajo régimen RESICO o Persona Física con Actividad Empresarial) "
        "cuyo importe antes de impuestos cuadre exactamente con el estado de cuenta generado por César."
    )

    # SECTION 6: MESA DE CORTE Y BONIFICACIONES POR DEFECTOS
    add_h1("6. Mesa de Corte: Bonificaciones por Defectos de Fábrica ('Banderas')")
    p = doc.add_paragraph(
        "Durante el proceso de corte de tela en almacén, es común que los rollos importados presenten defectos de origen "
        "(hasta 5 defectos por rollo: manchas, hilados corridos o marcas conocidas operativamente como 'banderas'). "
        "Para no demorar la entrega ni generar inconformidad con el cliente, el cortador aplica una bonificación física de metraje adicional (20 a 50 cm)."
    )

    add_callout(
        doc,
        "Las Dos Reglas de Oro en Bonificaciones de Corte",
        "1. EN FACTURA (COSTO CERO): La bonificación NUNCA debe facturarse ni sumarse al total a cobrar. Si se cortaron 1.50m (1.20m pedidos + 0.30m de compensación por bandera), la factura sale estrictamente por 1.20m. De lo contrario, se elevaría el saldo y se descuadraría el depósito bancario confirmado por César.\n\n"
        "2. EN INVENTARIO (DESCUENTO REAL DEL MARBETE): Los 0.30m regalados salieron físicamente del rollo matriz. Por ende, el cortador debe registrar la bonificación en su terminal para que el marbete HU remanente se actualice con el metraje real y no se generen 'metros fantasma' en el WMS.",
        bg_hex="FEF2F2",
        border_hex="EF4444"
    )

    # SECTION 7: RESERVAS PARA CLIENTES ESPECIALES
    add_h1("7. Reservas Protegidas para Clientes Especiales (Caso 'Seth')")
    p = doc.add_paragraph(
        "Existen clientes clave de alto volumen (ej. cliente 'Seth') que colocan pedidos superiores a 1,000 metros y cuyos ciclos "
        "corporativos de pago demoran de dos semanas a un mes. Para evitar que la regla estándar de liberación de apartados a las 72 horas "
        "cancele la reserva y provoque que otros agentes vendan la tela, el WMS 360° dispondrá de:\n"
        "• Un sub-almacén virtual / ubicación lógica: 'RESERVADO_CLIENTE_ESPECIAL'.\n"
        "• Los rollos quedan bloqueados para venta general, pero plenamente identificados para el cliente con un candado de tiempo autorizado por Dirección."
    )

    # SECTION 8: IMPACTO TÉCNICO EN SOFTWARE
    add_h1("8. Matriz de Cambios e Implementación en Software")
    
    tech_table = doc.add_table(rows=6, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    tech_headers = ["Componente / Archivo", "Ajuste Técnico Requerido", "Objetivo Operativo"]
    for i, h in enumerate(tech_headers):
        c = tech_table.rows[0].cells[i]
        set_cell_background(c, "0F172A")
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)

    tech_rows = [
        ("wms-backend/prisma/schema.prisma", "Campos retenidoCredito, bonificacionCorteMetros, diasCreditoAutorizados en ClientCreditConfig y Order", "Soporte de retención de crédito, bonificaciones a costo cero y plazos oficiales"),
        ("wms-backend/.../invoicing.service.ts", "Tolerancia de absorción de ±$10.00 MXN en registerPayment y bifurcación PUE (no timbrado) vs PPD (REP timbrado)", "Cuadre contable de César y cumplimiento fiscal con Facturapi"),
        ("wms-backend/.../commissions.service.ts", "Cálculo sobre bruto cobrado, deducción de notas de crédito y agrupación jerárquica (Vendedor, Gerente, Director)", "Reporte mensual día 10 para pago de comisiones con factura"),
        ("wms-frontend/.../CobranzaPage.tsx", "Pestaña 'Retenidos por Crédito' con botón de liberación discrecional por negociación de César", "Visibilidad y control total de pedidos retenidos por deuda vencida"),
        ("wms-frontend/.../CortadorView.tsx", "Campo de Bonificación por Defecto (metros) a costo cero que descuenta del rollo", "Eliminación de metros fantasma en almacén sin alterar cotización")
    ]

    for idx, (comp, req, obj) in enumerate(tech_rows):
        row = tech_table.rows[idx + 1]
        for c_idx, val in enumerate([comp, req, obj]):
            c = row.cells[c_idx]
            set_cell_background(c, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = INDIGO

    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(output_dir, "REGLAS_DE_NEGOCIO_COBRANZA_FACTURACION_FORMATEX.docx")
    doc.save(output_path)
    print(f"✅ Documento generado exitosamente en: {output_path}")

if __name__ == "__main__":
    build_cobranza_docx()
